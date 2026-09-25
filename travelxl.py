import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE = "travel.xlsx"
HEADERS = ["ID", "Name", "From", "To", "Date", "Transport", "Seats", "Price"]

# Create Excel
if not os.path.exists(FILE):
    wb = Workbook()
    wb.active.append(HEADERS)
    wb.save(FILE)

root = tk.Tk()
root.title("Travel Management System")
root.geometry("700x600")

entries = []


# ---------- Clear ----------
def clear():
    for e in entries:
        e.delete(0, tk.END)


# ---------- Add ----------
def add():
    data = [e.get() for e in entries]

    if "" in data:
        messagebox.showerror("Error", "Fill all fields")
        return

    wb = load_workbook(FILE)
    ws = wb.active
    ws.append(data)
    wb.save(FILE)

    messagebox.showinfo("Success", "Record Added")
    clear()


# ---------- View ----------
def view():
    win = tk.Toplevel(root)
    win.title("Travel Records")

    tree = ttk.Treeview(win, columns=HEADERS, show="headings")

    for h in HEADERS:
        tree.heading(h, text=h)
        tree.column(h, width=90)

    tree.pack(fill="both", expand=True)

    ws = load_workbook(FILE).active

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)


# ---------- Search ----------
def search():
    sid = entries[0].get()

    ws = load_workbook(FILE).active

    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == sid:
            messagebox.showinfo(
                "Found",
                "\n".join(f"{HEADERS[i]}: {row[i]}"
                          for i in range(len(HEADERS)))
            )
            return

    messagebox.showerror("Error", "Record Not Found")


# ---------- Update ----------
def update():
    sid = entries[0].get()
    data = [e.get() for e in entries]

    wb = load_workbook(FILE)
    ws = wb.active

    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value) == sid:

            for c in range(1, len(HEADERS) + 1):
                ws.cell(r, c).value = data[c - 1]

            wb.save(FILE)
            messagebox.showinfo("Success", "Record Updated")
            clear()
            return

    messagebox.showerror("Error", "Record Not Found")


# ---------- Delete ----------
def delete():
    sid = entries[0].get()

    if not messagebox.askyesno("Confirm", "Delete this record?"):
        return

    wb = load_workbook(FILE)
    ws = wb.active

    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value) == sid:
            ws.delete_rows(r)
            wb.save(FILE)
            messagebox.showinfo("Success", "Record Deleted")
            clear()
            return

    messagebox.showerror("Error", "Record Not Found")


# ---------- Dashboard ----------
def dashboard():
    for w in root.winfo_children():
        w.destroy()

    tk.Label(
        root,
        text="TRAVEL MANAGEMENT SYSTEM",
        font=("Arial", 20, "bold")
    ).pack(pady=15)

    global entries
    entries = []

    form = tk.Frame(root)
    form.pack()

    for i, h in enumerate(HEADERS):
        tk.Label(form, text=h).grid(row=i, column=0, pady=5)

        e = tk.Entry(form, width=30)
        e.grid(row=i, column=1)

        entries.append(e)

    buttons = [
        ("ADD", add),
        ("VIEW", view),
        ("SEARCH", search),
        ("UPDATE", update),
        ("DELETE", delete),
        ("CLEAR", clear),
        ("LOGOUT", login_page)
    ]

    for text, command in buttons:
        tk.Button(
            root,
            text=text,
            width=15,
            command=command
        ).pack(pady=3)


# ---------- Login ----------
def login_page():
    for w in root.winfo_children():
        w.destroy()

    tk.Label(
        root,
        text="TRAVEL MANAGEMENT SYSTEM",
        font=("Arial", 20, "bold")
    ).pack(pady=30)

    tk.Label(root, text="Username").pack()
    global user
    user = tk.Entry(root)
    user.pack()

    tk.Label(root, text="Password").pack()
    global pwd
    pwd = tk.Entry(root, show="*")
    pwd.pack()

    tk.Button(
        root,
        text="LOGIN",
        width=15,
        command=login
    ).pack(pady=20)


def login():
    if user.get() == "admin" and pwd.get() == "1234":
        dashboard()
    else:
        messagebox.showerror("Error", "Invalid Username or Password")


# Start
login_page()
root.mainloop()